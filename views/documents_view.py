# Path: views/documents_view.py
import sys
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, 
    QPushButton, QTableWidget, QTableWidgetItem, QHeaderView, 
    QFrame, QDialog, QFormLayout, QComboBox, QMessageBox, QTextEdit
)
from PyQt6.QtCore import Qt, QDate
from PyQt6.QtGui import QFont
import db  # Database handler module

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def show_alert(parent, title, message, icon=QMessageBox.Icon.Information, buttons=QMessageBox.StandardButton.Ok):
    """Custom styled alert popup to guarantee high-contrast, visible buttons."""
    msg = QMessageBox(parent)
    msg.setIcon(icon)
    msg.setWindowTitle(title)
    msg.setText(message)
    msg.setStandardButtons(buttons)
    msg.setStyleSheet("""
        QMessageBox { 
            background-color: #ffffff; 
        }
        QLabel { 
            color: #2c3e50; 
            font-size: 13px; 
            background-color: transparent;
        }
        QPushButton {
            background-color: #0b2545 !important;
            color: #ffffff !important;
            border: none !important;
            border-radius: 4px;
            padding: 6px 18px;
            font-weight: bold;
            min-width: 65px;
            min-height: 24px;
        }
        QPushButton:hover {
            background-color: #1a5b9e !important;
        }
        QPushButton:pressed {
            background-color: #05162d !important;
        }
    """)
    return msg.exec()


def format_publication_dates(issue_dates):
    """
    Formats a list of date objects/strings into standard affidavit format:
    e.g., [2026-09-28, 2026-10-05, 2026-10-12] -> "September 28, October 5, 12, 2026"
    """
    if not issue_dates:
        return ""

    parsed_dates = []
    for d in issue_dates:
        if isinstance(d, str):
            try:
                parsed_dates.append(datetime.strptime(d, "%Y-%m-%d").date())
            except ValueError:
                continue
        elif hasattr(d, "year"):
            parsed_dates.append(d)

    if not parsed_dates:
        return ""

    parsed_dates.sort()

    groups = []
    for d in parsed_dates:
        month_name = d.strftime("%B")
        year = d.year
        day = d.day

        if groups and groups[-1]["month"] == month_name and groups[-1]["year"] == year:
            groups[-1]["days"].append(day)
        else:
            groups.append({"month": month_name, "year": year, "days": [day]})

    parts = []
    last_year = parsed_dates[-1].year

    for group in groups:
        days_str = ", ".join(str(day) for day in group["days"])
        if group["year"] == last_year:
            parts.append(f"{group['month']} {days_str}")
        else:
            parts.append(f"{group['month']} {days_str}, {group['year']}")

    return f"{', '.join(parts)}, {last_year}"


class GenerateAffidavitDialog(QDialog):
    """Dialog to generate Affidavit of Publication."""
    def __init__(self, parent=None, record_data=None):
        super().__init__(parent)
        self.record_data = record_data
        self.setWindowTitle("Edit Affidavit Record" if record_data else "Generate Affidavit of Publication")
        self.setFixedSize(560, 540)
        self.setStyleSheet("QDialog { background-color: #f8f9fa; }")
        
        self.generated_file_path = ""
        self.publications_data = []

        self.init_ui()
        self.load_publications()

        if self.record_data:
            self.populate_existing_data()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(25, 25, 25, 25)
        layout.setSpacing(15)

        form = QFormLayout()
        form.setSpacing(12)

        input_style = """
            QLineEdit, QComboBox {
                padding: 8px 10px;
                border: 1px solid #ced4da;
                border-radius: 5px;
                background-color: white;
                color: #212529;
                font-size: 13px;
            }
            QLineEdit:focus, QComboBox:focus {
                border: 1px solid #0b2545;
            }
            QLineEdit:read-only {
                background-color: #e9ecef;
                color: #495057;
            }
        """

        textarea_style = """
            QTextEdit {
                padding: 8px 10px;
                border: 1px solid #ced4da;
                border-radius: 5px;
                background-color: white;
                color: #212529;
                font-size: 13px;
            }
            QTextEdit:focus {
                border: 1px solid #0b2545;
            }
        """

        # Select Publication
        self.pub_combo = QComboBox()
        self.pub_combo.setStyleSheet(input_style)
        self.pub_combo.currentIndexChanged.connect(self.on_publication_changed)

        # Affidavit Details
        self.affidavit_details_input = QTextEdit()
        self.affidavit_details_input.setPlaceholderText("Enter detailed affirmation statements here...")
        self.affidavit_details_input.setStyleSheet(textarea_style)
        self.affidavit_details_input.setFixedHeight(120)

        # Publication Dates
        self.pub_dates_input = QLineEdit()
        self.pub_dates_input.setPlaceholderText("e.g., September 28, October 5, 12, 2026")
        self.pub_dates_input.setStyleSheet(input_style)

        def make_label(text):
            lbl = QLabel(text)
            lbl.setFont(QFont("Arial", 9, QFont.Weight.Bold))
            lbl.setStyleSheet("color: #0b2545;")
            return lbl

        form.addRow(make_label("Select Publication:"), self.pub_combo)
        form.addRow(make_label("Affidavit Details:"), self.affidavit_details_input)
        form.addRow(make_label("Publication Dates:"), self.pub_dates_input)

        layout.addLayout(form)
        layout.addStretch()

        # Action Buttons
        btn_box = QHBoxLayout()
        btn_box.addStretch()

        cancel_btn = QPushButton("Cancel")
        cancel_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        cancel_btn.setStyleSheet("""
            QPushButton {
                padding: 8px 18px;
                border: 1px solid #ced4da;
                border-radius: 5px;
                background-color: white;
                color: #495057;
                font-weight: bold;
            }
            QPushButton:hover { background-color: #f1f3f5; }
        """)
        cancel_btn.clicked.connect(self.reject)

        generate_btn = QPushButton("Update Record" if self.record_data else "Generate and Save Affidavit")
        generate_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        generate_btn.setStyleSheet("""
            QPushButton {
                background-color: #0b2545;
                color: #ffffff;
                padding: 8px 20px;
                border-radius: 5px;
                font-weight: bold;
                border: none;
            }
            QPushButton:hover { background-color: #1a5b9e; }
        """)
        generate_btn.clicked.connect(self.generate_and_save_document)

        btn_box.addWidget(cancel_btn)
        btn_box.addWidget(generate_btn)

        layout.addLayout(btn_box)

    def load_publications(self):
        try:
            conn = db.get_connection()
            cursor = conn.cursor()
            
            cursor.execute("""
                SELECT p.publication_id, p.client_id, 
                       COALESCE(c.client_name, 'Unknown Client') AS client_name, 
                       COALESCE(p.case_number, '') AS case_no,
                       COALESCE(p.legal_description, '') AS legal_desc,
                       p.date_received
                FROM publication p
                LEFT JOIN client c ON p.client_id = c.client_id
                ORDER BY p.publication_id DESC
            """)

            rows = cursor.fetchall()
            self.pub_combo.clear()
            self.publications_data = []

            for row in rows:
                if isinstance(row, dict):
                    pub_id = row.get("publication_id")
                    client_id = row.get("client_id")
                    client_name = row.get("client_name") or "Unknown Client"
                    case_no = row.get("case_no") or ""
                    legal_desc = row.get("legal_desc") or ""
                    pub_date = row.get("date_received") or ""
                else:
                    pub_id = row[0]
                    client_id = row[1]
                    client_name = row[2] or "Unknown Client"
                    case_no = row[3] or ""
                    legal_desc = row[4] or ""
                    pub_date = row[5] or ""

                if not case_no and pub_id:
                    case_no = f"Pub #{pub_id}"

                # Fetch publication_issue dates for this publication
                cursor.execute("""
                    SELECT publication_date 
                    FROM publication_issue 
                    WHERE publication_id = %s 
                    ORDER BY publication_date ASC
                """, (pub_id,))
                
                issue_rows = cursor.fetchall()
                issue_dates = []
                for i_row in issue_rows:
                    if isinstance(i_row, dict):
                        issue_dates.append(i_row.get("publication_date"))
                    else:
                        issue_dates.append(i_row[0])

                formatted_dates_str = format_publication_dates(issue_dates)
                display_text = f"{case_no} - {client_name}"
                
                self.publications_data.append({
                    "publication_id": pub_id,
                    "client_id": client_id,
                    "client_name": client_name,
                    "case_no": case_no,
                    "legal_description": legal_desc,
                    "publication_date": str(pub_date),
                    "formatted_issue_dates": formatted_dates_str
                })
                self.pub_combo.addItem(display_text)

            cursor.close()
            conn.close()

            if self.publications_data and not self.record_data:
                self.on_publication_changed(0)

        except Exception as e:
            print(f"Error loading publications: {e}")
            self.pub_combo.addItem("No publications found")

    def on_publication_changed(self, index):
        if 0 <= index < len(self.publications_data):
            pub_info = self.publications_data[index]
            legal_desc = pub_info.get("legal_description", "").strip()
            client_name = pub_info.get("client_name", "").strip()
            case_no = pub_info.get("case_no", "").strip()

            formatted_text = f"{legal_desc}\n\n{client_name}\nCase No. {case_no}".strip()
            self.affidavit_details_input.setPlainText(formatted_text)

            formatted_dates = pub_info.get("formatted_issue_dates", "")
            self.pub_dates_input.setText(formatted_dates)

    def populate_existing_data(self):
        existing_case = str(self.record_data.get("case_no", ""))
        for i in range(self.pub_combo.count()):
            if existing_case in self.pub_combo.itemText(i):
                self.pub_combo.setCurrentIndex(i)
                break

        details = self.record_data.get("details", "")
        if details:
            self.affidavit_details_input.setPlainText(details)

    def generate_docx_file(self, client_name, case_no, pub_dates, affidavit_details):
        try:
            template_path = os.path.join(BASE_DIR, "assets", "Affidavit template.docx")
            if not os.path.exists(template_path):
                if os.path.exists(os.path.join(BASE_DIR, "assets", "Affidavit template_7.docx")):
                    template_path = os.path.join(BASE_DIR, "assets", "Affidavit template_7.docx")
                elif os.path.exists(os.path.join(BASE_DIR, "assets", "Affidavit template_6.docx")):
                    template_path = os.path.join(BASE_DIR, "assets", "Affidavit template_6.docx")

            if os.path.exists(template_path):
                doc = Document(template_path)
            else:
                doc = Document()

            today = QDate.currentDate()
            day = today.day()
            
            if 11 <= day <= 13:
                suffix = 'th'
            else:
                suffix = {1: 'st', 2: 'nd', 3: 'rd'}.get(day % 10, 'th')
            
            signing_date_str = f"{day}{suffix} day of {today.toString('MMMM, yyyy')}"

            placeholders = {
                "(Affidavit Details)": affidavit_details,
                "(Publication Dates)": pub_dates,
                "(Current Date the Affidavit was made)": signing_date_str
            }

            def replace_in_paragraphs(paragraphs):
                for p in paragraphs:
                    for placeholder, value in placeholders.items():
                        if placeholder in p.text:
                            inline = p.runs
                            full_text = ''.join(run.text for run in inline)
                            if placeholder in full_text:
                                full_text = full_text.replace(placeholder, value)
                                for i, run in enumerate(inline):
                                    if i == 0:
                                        run.text = full_text
                                    else:
                                        run.text = ""

            replace_in_paragraphs(doc.paragraphs)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        replace_in_paragraphs(cell.paragraphs)

            os.makedirs("generated_documents", exist_ok=True)
            safe_client = "".join(c for c in client_name if c.isalnum() or c.isspace()).strip().replace(" ", "_")
            filename = f"Affidavit_{safe_client}_{case_no.replace('/', '_')}.docx"
            file_path = os.path.abspath(os.path.join("generated_documents", filename))
            
            doc.save(file_path)
            return file_path
        except Exception as e:
            print(f"Error generating docx: {e}")
            return ""

    def generate_and_save_document(self):
        curr_index = self.pub_combo.currentIndex()
        if curr_index < 0 or curr_index >= len(self.publications_data):
            show_alert(self, "Validation Error", "Please select a valid publication record.", QMessageBox.Icon.Warning)
            return

        pub_info = self.publications_data[curr_index]
        client_name = pub_info.get("client_name", "Client")
        client_id = pub_info.get("client_id")
        case_no = pub_info.get("case_no", "")
        doc_type = "Affidavit of Publication"
        pub_dates = self.pub_dates_input.text().strip()
        affidavit_details = self.affidavit_details_input.toPlainText().strip()
        date_uploaded = QDate.currentDate().toString("yyyy-MM-dd")

        if not case_no:
            show_alert(self, "Validation Error", "Selected publication has no Case Number.", QMessageBox.Icon.Warning)
            return

        if not pub_dates:
            show_alert(self, "Validation Error", "Please enter publication dates.", QMessageBox.Icon.Warning)
            return

        file_path = self.generate_docx_file(client_name, case_no, pub_dates, affidavit_details)
        if not file_path:
            show_alert(self, "Generation Error", "Failed to generate Word document file.", QMessageBox.Icon.Critical)
            return

        try:
            conn = db.get_connection()
            cursor = conn.cursor()

            if self.record_data:
                doc_id = self.record_data.get("document_id")
                query = """
                    UPDATE document 
                    SET case_no = %s, client_id = %s, client_name = %s, document_type = %s, file_path = %s
                    WHERE document_id = %s
                """
                cursor.execute(query, (case_no, client_id, client_name, doc_type, file_path, doc_id))
                msg = "Affidavit record updated successfully!"
            else:
                query = """
                    INSERT INTO document (case_no, client_id, client_name, document_type, file_path, date_uploaded, status)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                """
                cursor.execute(query, (case_no, client_id, client_name, doc_type, file_path, date_uploaded, "Generated"))
                msg = "Affidavit generated and saved successfully!"

            conn.commit()
            cursor.close()
            conn.close()

            show_alert(self, "Success", f"{msg}\n\nSaved to:\n{file_path}", QMessageBox.Icon.Information)
            self.accept()
        except Exception as e:
            show_alert(self, "Database Error", f"Failed to save affidavit record:\n{str(e)}", QMessageBox.Icon.Critical)


class GenerateBillingDialog(QDialog):
    """Dialog to generate Billing Statement using Billing Statement Template.docx."""
    def __init__(self, parent=None, record_data=None):
        super().__init__(parent)
        self.record_data = record_data
        self.setWindowTitle("Edit Billing Record" if record_data else "Generate Billing Statement")
        self.setFixedSize(560, 540)
        self.setStyleSheet("QDialog { background-color: #f8f9fa; }")
        
        self.generated_file_path = ""
        self.publications_data = []

        self.init_ui()
        self.load_publications()

        if self.record_data:
            self.populate_existing_data()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(25, 25, 25, 25)
        layout.setSpacing(15)

        form = QFormLayout()
        form.setSpacing(12)

        input_style = """
            QLineEdit, QComboBox {
                padding: 8px 10px;
                border: 1px solid #ced4da;
                border-radius: 5px;
                background-color: white;
                color: #212529;
                font-size: 13px;
            }
            QLineEdit:focus, QComboBox:focus {
                border: 1px solid #0b2545;
            }
        """

        textarea_style = """
            QTextEdit {
                padding: 8px 10px;
                border: 1px solid #ced4da;
                border-radius: 5px;
                background-color: white;
                color: #212529;
                font-size: 13px;
            }
            QTextEdit:focus {
                border: 1px solid #0b2545;
            }
        """

        # Select Publication
        self.pub_combo = QComboBox()
        self.pub_combo.setStyleSheet(input_style)
        self.pub_combo.currentIndexChanged.connect(self.on_publication_changed)

        # Billing Details (Particulars)
        self.billing_details_input = QTextEdit()
        self.billing_details_input.setPlaceholderText("Enter billing particulars / legal description here...")
        self.billing_details_input.setStyleSheet(textarea_style)
        self.billing_details_input.setFixedHeight(120)

        # Billing Amount
        self.amount_input = QLineEdit()
        self.amount_input.setPlaceholderText("e.g., 4518.00")
        self.amount_input.setStyleSheet(input_style)

        def make_label(text):
            lbl = QLabel(text)
            lbl.setFont(QFont("Arial", 9, QFont.Weight.Bold))
            lbl.setStyleSheet("color: #0b2545;")
            return lbl

        form.addRow(make_label("Select Publication:"), self.pub_combo)
        form.addRow(make_label("Billing Details:"), self.billing_details_input)
        form.addRow(make_label("Billing Amount (PHP):"), self.amount_input)

        layout.addLayout(form)
        layout.addStretch()

        # Action Buttons
        btn_box = QHBoxLayout()
        btn_box.addStretch()

        cancel_btn = QPushButton("Cancel")
        cancel_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        cancel_btn.setStyleSheet("""
            QPushButton {
                padding: 8px 18px;
                border: 1px solid #ced4da;
                border-radius: 5px;
                background-color: white;
                color: #495057;
                font-weight: bold;
            }
            QPushButton:hover { background-color: #f1f3f5; }
        """)
        cancel_btn.clicked.connect(self.reject)

        generate_btn = QPushButton("Update Billing" if self.record_data else "Generate and Save Billing")
        generate_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        generate_btn.setStyleSheet("""
            QPushButton {
                background-color: #0b2545;
                color: #ffffff;
                padding: 8px 20px;
                border-radius: 5px;
                font-weight: bold;
                border: none;
            }
            QPushButton:hover { background-color: #1a5b9e; }
        """)
        generate_btn.clicked.connect(self.generate_and_save_billing)

        btn_box.addWidget(cancel_btn)
        btn_box.addWidget(generate_btn)
        layout.addLayout(btn_box)

    def load_publications(self):
        try:
            conn = db.get_connection()
            cursor = conn.cursor()
            cursor.execute("""
                SELECT p.publication_id, p.client_id, 
                       COALESCE(c.client_name, 'Unknown Client') AS client_name, 
                       COALESCE(c.address, '') AS address,
                       COALESCE(p.case_number, '') AS case_no,
                       COALESCE(p.legal_description, '') AS legal_description,
                       COALESCE(p.total_amount, 3000.00) AS total_amount
                FROM publication p
                LEFT JOIN client c ON p.client_id = c.client_id
                ORDER BY p.publication_id DESC
            """)
            rows = cursor.fetchall()
            self.pub_combo.clear()
            self.publications_data = []

            for row in rows:
                if isinstance(row, dict):
                    pub_id = row.get("publication_id")
                    client_id = row.get("client_id")
                    client_name = row.get("client_name") or "Unknown Client"
                    address = row.get("address") or ""
                    case_no = row.get("case_no") or ""
                    legal_desc = row.get("legal_description") or ""
                    total_amt = row.get("total_amount") or 3000.00
                else:
                    pub_id = row[0]
                    client_id = row[1]
                    client_name = row[2] or "Unknown Client"
                    address = row[3] or ""
                    case_no = row[4] or ""
                    legal_desc = row[5] or ""
                    total_amt = row[6] or 3000.00

                if not case_no and pub_id:
                    case_no = f"Pub #{pub_id}"

                display_text = f"{case_no} - {client_name}"
                self.publications_data.append({
                    "publication_id": pub_id,
                    "client_id": client_id,
                    "client_name": client_name,
                    "address": address,
                    "case_no": case_no,
                    "legal_description": legal_desc,
                    "total_amount": float(total_amt)
                })
                self.pub_combo.addItem(display_text)
            cursor.close()
            conn.close()

            if self.publications_data and not self.record_data:
                self.on_publication_changed(0)

        except Exception as e:
            print(f"Error loading publications for billing: {e}")

    def on_publication_changed(self, index):
        if 0 <= index < len(self.publications_data):
            pub_info = self.publications_data[index]
            self.amount_input.setText(str(pub_info.get("total_amount", 3000.00)))

            legal_desc = pub_info.get("legal_description", "").strip()
            client_name = pub_info.get("client_name", "").strip()
            case_no = pub_info.get("case_no", "").strip()

            formatted_text = f"{legal_desc}\n\n{client_name}\nCase No. {case_no}".strip()
            self.billing_details_input.setPlainText(formatted_text)

    def populate_existing_data(self):
        existing_case = str(self.record_data.get("case_no", ""))
        for i in range(self.pub_combo.count()):
            if existing_case in self.pub_combo.itemText(i):
                self.pub_combo.setCurrentIndex(i)
                break

    def number_to_words(self, n):
        if n == 0:
            return "ZERO PESOS"
        ones = ["", "ONE", "TWO", "THREE", "FOUR", "FIVE", "SIX", "SEVEN", "EIGHT", "NINE", 
                "TEN", "ELEVEN", "TWELVE", "THIRTEEN", "FOURTEEN", "FIFTEEN", "SIXTEEN", 
                "SEVENTEEN", "EIGHTEEN", "NINETEEN"]
        tens = ["", "", "TWENTY", "THIRTY", "FORTY", "FIFTY", "SIXTY", "SEVENTY", "EIGHTY", "NINETY"]

        def _to_words(num):
            if num < 20:
                return ones[num]
            elif num < 100:
                return tens[num // 10] + (" " + ones[num % 10] if num % 10 != 0 else "")
            elif num < 1000:
                return ones[num // 100] + " HUNDRED" + (" " + _to_words(num % 100) if num % 100 != 0 else "")
            elif num < 1000000:
                thousands = num // 1000
                remainder = num % 1000
                res = _to_words(thousands) + " THOUSAND"
                if remainder > 0:
                    res += " " + _to_words(remainder)
                return res
            return str(num)

        n_int = int(n)
        n_cents = int(round((n - n_int) * 100))
        words = _to_words(n_int) + " PESOS"
        if n_cents > 0:
            words += f" AND {n_cents:02d}/100"
        return words

    def generate_docx_file(self, client_name, client_address, case_no, billing_details, amount):
        try:
            template_path = os.path.join(BASE_DIR, "assets", "Billing Statement Template.docx")
            if not os.path.exists(template_path):
                if os.path.exists(os.path.join(BASE_DIR, "assets", "Billing Statement Template_2.docx")):
                    template_path = os.path.join(BASE_DIR, "assets", "Billing Statement Template_2.docx")

            if os.path.exists(template_path):
                doc = Document(template_path)
            else:
                doc = Document()

            current_date_str = QDate.currentDate().toString("MMMM d, yyyy")
            amount_float = float(amount)
            amount_in_words = self.number_to_words(amount_float)
            formatted_amount = f"P{amount_float:,.2f}"

            placeholders = {
                "(Current Date Billing Statement was Made)": current_date_str,
                "(Client Name)": client_name,
                "(Adress)": client_address,
                "(Address)": client_address,
                "(Billing Details)": billing_details if billing_details else f"Publication fee for Case No. {case_no}",
                "(Amount Total Worded Format)": amount_in_words,
                "(Amount Total)": formatted_amount
            }

            def replace_in_paragraphs(paragraphs):
                for p in paragraphs:
                    for placeholder, value in placeholders.items():
                        if placeholder in p.text:
                            inline = p.runs
                            full_text = ''.join(run.text for run in inline)
                            if placeholder in full_text:
                                full_text = full_text.replace(placeholder, value)
                                for i, run in enumerate(inline):
                                    if i == 0:
                                        run.text = full_text
                                    else:
                                        run.text = ""

            replace_in_paragraphs(doc.paragraphs)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        replace_in_paragraphs(cell.paragraphs)

            os.makedirs("generated_documents", exist_ok=True)
            safe_client = "".join(c for c in client_name if c.isalnum() or c.isspace()).strip().replace(" ", "_")
            filename = f"Billing_{safe_client}_{case_no.replace('/', '_')}.docx"
            file_path = os.path.abspath(os.path.join("generated_documents", filename))
            
            doc.save(file_path)
            return file_path
        except Exception as e:
            print(f"Error generating billing docx: {e}")
            return ""

    def generate_and_save_billing(self):
        curr_index = self.pub_combo.currentIndex()
        if curr_index < 0 or curr_index >= len(self.publications_data):
            show_alert(self, "Validation Error", "Please select a valid publication record.", QMessageBox.Icon.Warning)
            return

        pub_info = self.publications_data[curr_index]
        client_name = pub_info.get("client_name", "Client")
        client_id = pub_info.get("client_id")
        client_address = pub_info.get("address", "")
        case_no = pub_info.get("case_no", "")
        billing_details = self.billing_details_input.toPlainText().strip()
        doc_type = "Billing Statement"

        try:
            amount = float(self.amount_input.text().strip() or "0.00")
        except ValueError:
            show_alert(self, "Validation Error", "Please enter a valid numeric billing amount.", QMessageBox.Icon.Warning)
            return

        if not case_no:
            show_alert(self, "Validation Error", "Selected publication has no Case Number.", QMessageBox.Icon.Warning)
            return

        file_path = self.generate_docx_file(client_name, client_address, case_no, billing_details, amount)
        if not file_path:
            show_alert(self, "Generation Error", "Failed to generate Billing Word document.", QMessageBox.Icon.Critical)
            return

        try:
            conn = db.get_connection()
            if not conn:
                show_alert(self, "Database Error", "Failed to establish a database connection.", QMessageBox.Icon.Critical)
                return

            cursor = conn.cursor()

            if self.record_data:
                doc_id = self.record_data.get("document_id")
                query = """
                    UPDATE document 
                    SET case_no = %s, client_id = %s, client_name = %s, document_type = %s, file_path = %s
                    WHERE document_id = %s
                """
                cursor.execute(query, (case_no, client_id, client_name, doc_type, file_path, doc_id))
                msg = "Billing record updated successfully!"
            else:
                query = """
                    INSERT INTO document (case_no, client_id, client_name, document_type, file_path, date_uploaded, status)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                """
                cursor.execute(query, (case_no, client_id, client_name, doc_type, file_path, QDate.currentDate().toString("yyyy-MM-dd"), "Generated"))
                msg = "Billing Statement generated and saved successfully!"

            cursor.execute("UPDATE publication SET total_amount = %s WHERE publication_id = %s", (amount, pub_info.get("publication_id")))

            if hasattr(conn, 'commit'):
                conn.commit()

            cursor.close()
            if hasattr(conn, 'close'):
                conn.close()

            show_alert(self, "Success", f"{msg}\n\nSaved to:\n{file_path}", QMessageBox.Icon.Information)
            self.accept()
        except Exception as e:
            show_alert(self, "Database Error", f"Failed to save billing record:\n{str(e)}", QMessageBox.Icon.Critical)


class DocumentsView(QWidget):
    def __init__(self):
        super().__init__()
        self.init_ui()
        self.load_documents()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(25, 25, 25, 25)
        layout.setSpacing(20)

        header_layout = QHBoxLayout()

        title = QLabel("Documents")
        title.setFont(QFont("Arial", 16, QFont.Weight.Bold))
        title.setStyleSheet("color: #0b2545;")

        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Search by case number or document type...")
        self.search_input.setFixedWidth(320)
        self.search_input.setStyleSheet("""
            QLineEdit {
                padding: 8px 12px;
                border: 1px solid #dcdfe6;
                border-radius: 6px;
                background-color: white;
                color: #2c3e50;
            }
            QLineEdit:focus { border: 1px solid #0b2545; }
        """)
        self.search_input.textChanged.connect(self.filter_documents)

        gen_affidavit_btn = QPushButton("Generate Affidavit")
        gen_affidavit_btn.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        gen_affidavit_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        gen_affidavit_btn.setStyleSheet("""
            QPushButton {
                background-color: #0b2545;
                color: white;
                padding: 9px 16px;
                border-radius: 6px;
                border: none;
            }
            QPushButton:hover { background-color: #1a5b9e; }
        """)
        gen_affidavit_btn.clicked.connect(self.open_generate_affidavit_dialog)

        gen_billing_btn = QPushButton("Generate Billing")
        gen_billing_btn.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        gen_billing_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        gen_billing_btn.setStyleSheet("""
            QPushButton {
                background-color: #0b2545;
                color: white;
                padding: 9px 16px;
                border-radius: 6px;
                border: none;
            }
            QPushButton:hover { background-color: #1a5b9e; }
        """)
        gen_billing_btn.clicked.connect(self.open_generate_billing_dialog)

        header_layout.addWidget(title)
        header_layout.addSpacing(20)
        header_layout.addWidget(self.search_input)
        header_layout.addStretch()
        header_layout.addWidget(gen_affidavit_btn)
        header_layout.addWidget(gen_billing_btn)

        layout.addLayout(header_layout)

        card_frame = QFrame()
        card_frame.setStyleSheet("QFrame { background-color: white; border-radius: 8px; border: 1px solid #e0e0e0; }")
        card_layout = QVBoxLayout(card_frame)
        card_layout.setContentsMargins(15, 15, 15, 15)

        self.table = QTableWidget()
        self.table.setColumnCount(5)
        self.table.setHorizontalHeaderLabels([
            "No.", "Case No.", "Client", "Document Type", "Actions"
        ])
        
        self.table.setStyleSheet("""
            QTableWidget {
                border: none;
                gridline-color: #f0f0f0;
                selection-background-color: #f4f6f9;
                selection-color: #2c3e50;
            }
            QHeaderView::section {
                background-color: #ffffff;
                color: #2c3e50;
                font-weight: bold;
                padding: 10px;
                border: none;
                border-bottom: 2px solid #e0e0e0;
            }
            QTableWidget::item {
                padding: 8px;
                border-bottom: 1px solid #f0f0f0;
                color: #2c3e50;
            }
        """)

        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(4, QHeaderView.ResizeMode.ResizeToContents)

        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(52)
        self.table.setAlternatingRowColors(True)

        card_layout.addWidget(self.table)
        layout.addWidget(card_frame)

    def load_documents(self):
        try:
            conn = db.get_connection()
            cursor = conn.cursor()

            query = """
                SELECT document_id, case_no, client_id, client_name, document_type, file_path, date_uploaded, status 
                FROM document 
                ORDER BY document_id DESC
            """
            cursor.execute(query)
            records = cursor.fetchall()

            cursor.close()
            conn.close()

            self.populate_table(records)

        except Exception as e:
            show_alert(self, "Database Error", f"Failed to load document records:\n{str(e)}", QMessageBox.Icon.Critical)

    def populate_table(self, records):
        self.table.setRowCount(0)

        for index, row in enumerate(records):
            if isinstance(row, dict):
                d_id = row.get("document_id")
                case_no = str(row.get("case_no", ""))
                client_id = row.get("client_id")
                client_name = str(row.get("client_name", ""))
                doc_type = str(row.get("document_type", ""))
                file_path = str(row.get("file_path", ""))
                date_up = str(row.get("date_uploaded", ""))
                status = str(row.get("status", ""))
            else:
                d_id, case_no, client_id, client_name, doc_type, file_path, date_up, status = row

            row_position = self.table.rowCount()
            self.table.insertRow(row_position)
            self.table.setRowHeight(row_position, 52)

            no_item = QTableWidgetItem(str(index + 1))
            no_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)

            case_item = QTableWidgetItem(case_no)
            client_item = QTableWidgetItem(client_name)
            doc_type_item = QTableWidgetItem(doc_type)

            self.table.setItem(row_position, 0, no_item)
            self.table.setItem(row_position, 1, case_item)
            self.table.setItem(row_position, 2, client_item)
            self.table.setItem(row_position, 3, doc_type_item)

            actions_widget = QWidget()
            actions_layout = QHBoxLayout(actions_widget)
            actions_layout.setContentsMargins(6, 4, 6, 4)
            actions_layout.setSpacing(10)
            actions_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)

            open_btn = QPushButton("Open")
            open_btn.setCursor(Qt.CursorShape.PointingHandCursor)
            open_btn.setMinimumWidth(65)
            open_btn.setMinimumHeight(30)
            open_btn.setStyleSheet("""
                QPushButton {
                    background-color: #0b2545;
                    color: white;
                    border: none;
                    border-radius: 4px;
                    padding: 6px 12px;
                    font-weight: bold;
                    font-size: 12px;
                }
                QPushButton:hover { background-color: #1a5b9e; }
            """)
            open_btn.clicked.connect(lambda checked, path=file_path: self.open_file(path))

            delete_btn = QPushButton("Delete")
            delete_btn.setCursor(Qt.CursorShape.PointingHandCursor)
            delete_btn.setMinimumWidth(65)
            delete_btn.setMinimumHeight(30)
            delete_btn.setStyleSheet("""
                QPushButton {
                    background-color: #e74c3c;
                    color: white;
                    border: none;
                    border-radius: 4px;
                    padding: 6px 12px;
                    font-weight: bold;
                    font-size: 12px;
                }
                QPushButton:hover { background-color: #c0392b; }
            """)
            delete_btn.clicked.connect(lambda checked, id_val=d_id: self.delete_document(id_val))

            actions_layout.addWidget(open_btn)
            actions_layout.addWidget(delete_btn)

            self.table.setCellWidget(row_position, 4, actions_widget)

    def open_file(self, file_path):
        if file_path and os.path.exists(file_path):
            try:
                os.startfile(file_path)
            except Exception as e:
                show_alert(self, "Error", f"Could not open file:\n{str(e)}", QMessageBox.Icon.Critical)
        else:
            show_alert(self, "File Not Found", "The generated document file could not be located on disk.", QMessageBox.Icon.Warning)

    def filter_documents(self):
        search_text = self.search_input.text().lower()
        for row in range(self.table.rowCount()):
            case_no = self.table.item(row, 1).text().lower() if self.table.item(row, 1) else ""
            client = self.table.item(row, 2).text().lower() if self.table.item(row, 2) else ""
            doc_type = self.table.item(row, 3).text().lower() if self.table.item(row, 3) else ""

            match = search_text in case_no or search_text in client or search_text in doc_type
            self.table.setRowHidden(row, not match)

    def open_generate_affidavit_dialog(self):
        dialog = GenerateAffidavitDialog(self)
        if dialog.exec():
            self.load_documents()

    def open_generate_billing_dialog(self):
        dialog = GenerateBillingDialog(self)
        if dialog.exec():
            self.load_documents()

    def open_edit_dialog(self, record_data):
        if record_data.get("document_type") == "Billing Statement":
            dialog = GenerateBillingDialog(self, record_data=record_data)
        else:
            dialog = GenerateAffidavitDialog(self, record_data=record_data)
        if dialog.exec():
            self.load_documents()

    def delete_document(self, doc_id):
        confirm = show_alert(
            self, 
            "Confirm Delete", 
            "Are you sure you want to delete this document record?", 
            QMessageBox.Icon.Question, 
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if confirm == QMessageBox.StandardButton.Yes:
            try:
                conn = db.get_connection()
                cursor = conn.cursor()
                cursor.execute("DELETE FROM document WHERE document_id = %s", (doc_id,))
                conn.commit()
                cursor.close()
                conn.close()
                self.load_documents()
            except Exception as e:
                show_alert(self, "Database Error", f"Failed to delete document record:\n{str(e)}", QMessageBox.Icon.Critical)